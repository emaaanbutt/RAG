from pathlib import Path

from langchain_core.documents import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from rag.schema import Measurement

WHO_COLUMNS = {
    "IND_NAME",
    "DIM_GEO_NAME",
    "DIM_TIME_YEAR",
    "DIM_1_CODE",
    "VALUE_NUMERIC",
}


class ExcelLoader:
    def __init__(self, project_path : Path):
        self.project_path = project_path


    def load(self, path: Path) -> tuple[list[Document], list[Measurement]]:

        source = str(path.relative_to(self.project_path))
        documents = []
        measurements = []

        work_book = load_workbook(
            path,
            read_only=True,
            data_only=True
        )

        try:
            for sheet in work_book:
                rows = sheet.iter_rows(values_only=True)
                first_row = next(rows, None)

                if first_row is None:
                    continue

                headers = []

                for i, value in enumerate(first_row, start=1):
                    if value is not None:
                        headers.append(str(value).strip())

                    else:
                        headers.append(f"column_{i}")

                if sheet.max_row == 1:
                    documents.append(Document(
                        page_content=" | ".join(
                            str(value)
                            for value in first_row
                            if value is not None
                        ),
                        metadata={
                            "source": source,
                            "sheet": sheet.title,
                            "cell_range": "A1",
                            "row": 1,
                            "kind": "xlsx_note",
                        },
                    ))

                    continue

                for row_number, values in enumerate(rows, start=2):
                    pairs = []

                    for i, value in enumerate(values):
                        if values is not None and str(value).strip():
                            pairs.append((headers[i], value))

                    if not pairs:
                        continue

                    fields = dict(pairs)
                    text = " | ".join(
                        f"{name}: {value}"
                        for name, value in pairs
                    )
                    last_column = get_column_letter(len(values))

                    metadata = {
                        "source": source,
                        "sheet": sheet.title,
                        "row": row_number,
                        "cell_range": (
                            f"A{row_number}:"
                            f"{last_column}{row_number}"
                        ),
                        "kind": "xlsx_row",
                    }

                    documents.append(Document(
                        page_content=text,
                        metadata=metadata,
                    ))

                    if (
                        WHO_COLUMNS.issubset(headers)
                        and fields.get("VALUE_NUMERIC") is not None
                    ):
                        measurements.append(Measurement(
                            indicator=str(fields["IND_NAME"]),
                            geography=str(fields["DIM_GEO_NAME"]),
                            year=int(fields["DIM_TIME_YEAR"]),
                            dimension=str(fields["DIM_1_CODE"]),
                            value=float(fields["VALUE_NUMERIC"]),
                            displayed=str(
                                fields.get("VALUE_STRING", "")
                            ),
                            comment=str(
                                fields.get("VALUE_COMMENTS") or ""
                            ),
                            source=source,
                            sheet=sheet.title,
                            row=row_number,
                        ))
        finally:
            work_book.close()

        return documents, measurements